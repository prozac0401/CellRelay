"""Thread-local Excel automation for files that require the desktop application.

Only plain Python values leave these functions. Each operation owns its Excel
instance and closes it before returning, so readers do not lock later writes or
pass COM objects between the UI and the Excel worker thread.
"""

from __future__ import annotations

import gc
import logging
import sys
import time
import traceback
from contextlib import contextmanager
from functools import partial
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)
SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm", ".xls", ".xlsb"}
EXCEL_FILE_FILTER = "Excel Workbook (*.xlsx *.xlsm *.xls *.xlsb)"
_READ_BATCH_SIZE = 512
_COM_RETRY_TIMEOUT = 3.0
_COM_RETRY_MAX_ATTEMPTS = 6
_COM_RETRYABLE_HRESULTS = {0x80010001, 0x8001010A}


def _com_call(operation, callback, *, retry_rejected=True):
    """Retry only calls Excel explicitly rejected before executing them.

    This bounds retry delays, not a COM call that never returns. A hard timeout
    for such calls requires a separately supervised process.
    """
    started = time.monotonic()
    logger.info("Excel COM started operation=%s", operation)
    for attempt in range(1, _COM_RETRY_MAX_ATTEMPTS + 1):
        try:
            result = callback()
        except Exception as exc:
            elapsed = time.monotonic() - started
            hresult = getattr(exc, "hresult", None)
            code = hresult & 0xFFFFFFFF if isinstance(hresult, int) else None
            code_text = f"0x{code:08X}" if code is not None else "unavailable"
            remaining = _COM_RETRY_TIMEOUT - elapsed
            if (
                retry_rejected
                and code in _COM_RETRYABLE_HRESULTS
                and attempt < _COM_RETRY_MAX_ATTEMPTS
                and remaining > 0
            ):
                delay = min(0.1 * 2 ** (attempt - 1), 1.0, remaining)
                logger.warning(
                    "Excel COM rejected operation=%s attempt=%d elapsed=%.3fs "
                    "hresult=%s retry_in=%.3fs",
                    operation,
                    attempt,
                    elapsed,
                    code_text,
                    delay,
                )
                time.sleep(delay)
                if time.monotonic() - started < _COM_RETRY_TIMEOUT:
                    continue
            logger.exception(
                "Excel COM failed operation=%s attempts=%d elapsed=%.3fs hresult=%s",
                operation,
                attempt,
                time.monotonic() - started,
                code_text,
            )
            # Failed callback/pywin32 frames can retain cell proxies until after
            # Quit, even when callers clear their references in finally. Keep
            # the traceback locations but release locals in unwound frames;
            # clear_frames leaves this still-executing frame alone.
            traceback.clear_frames(exc.__traceback__)
            callback = None
            raise
        logger.info(
            "Excel COM completed operation=%s attempts=%d elapsed=%.3fs",
            operation,
            attempt,
            time.monotonic() - started,
        )
        return result


def is_ooxml_workbook(path: Path) -> bool:
    # Legacy OLE files can contain an embedded theme ZIP; is_zipfile alone
    # would misidentify that inner package as the workbook itself.
    with path.open("rb") as stream:
        if stream.read(4) != b"PK\x03\x04":
            return False
    try:
        with ZipFile(path) as package:
            return "xl/workbook.xml" in package.namelist()
    except BadZipFile:
        return False


class ExcelComUnavailable(RuntimeError):
    """The Windows desktop Excel automation runtime is unavailable."""


def _load_runtime():
    if sys.platform != "win32":
        raise ExcelComUnavailable("Excel COM은 Windows에서만 사용할 수 있습니다.")
    try:
        import pythoncom
        import win32com.client
    except ImportError as exc:
        raise ExcelComUnavailable(
            "Excel COM 구성요소(pywin32)가 없습니다. "
            "requirements.txt를 다시 설치하거나 최신 배포 EXE를 사용하세요."
        ) from exc
    return pythoncom, win32com.client


@contextmanager
def excel_application():
    pythoncom, client = _load_runtime()
    pythoncom.CoInitializeEx(pythoncom.COINIT_APARTMENTTHREADED)
    application = None
    try:
        try:
            # Never attach to or quit the user's existing Excel instance.
            application = _com_call(
                "DispatchEx",
                lambda: client.DispatchEx("Excel.Application"),
                retry_rejected=False,
            )
        except Exception as exc:
            raise ExcelComUnavailable(
                "Microsoft Excel COM을 시작하지 못했습니다. "
                "Windows용 데스크톱 Excel 설치 및 실행 상태를 확인하세요."
            ) from exc
        for name, value in (
            ("Visible", False),
            ("DisplayAlerts", False),
            ("EnableEvents", False),
            ("AskToUpdateLinks", False),
            ("AutomationSecurity", 3),  # msoAutomationSecurityForceDisable
        ):
            _com_call(
                f"Application.{name}.set",
                lambda name=name, value=value: setattr(application, name, value),
            )
        yield application
    finally:
        # Release worksheet/type-info wrappers while their Excel server is
        # still alive. Callers clear their own app/book references in finally.
        gc.collect()
        if application is not None:
            try:
                _com_call("Application.Quit", lambda: application.Quit())
            except Exception:
                logger.exception("Could not quit the CellRelay Excel instance")
        application = None
        gc.collect()
        pythoncom.CoUninitialize()


@contextmanager
def open_workbook(application, path: Path, *, read_only: bool):
    workbook = None
    try:
        try:
            workbook = _com_call(
                "Workbooks.Open",
                lambda: application.Workbooks.Open(
                    Filename=str(path),
                    UpdateLinks=0,
                    ReadOnly=read_only,
                    # Explicit empty passwords avoid invisible password dialogs.
                    Password="",
                    WriteResPassword="",
                    IgnoreReadOnlyRecommended=True,
                    Notify=False,
                    AddToMru=False,
                ),
            )
        except Exception as exc:
            raise RuntimeError(
                "Excel COM으로 파일을 열지 못했습니다. "
                "같은 PC의 Excel에서 파일이 열리는지, 암호·보안 승인이나 "
                f"접근 권한이 필요한지 확인하세요: {exc}"
            ) from exc
        if not read_only and _com_call("Workbook.ReadOnly", lambda: workbook.ReadOnly):
            raise PermissionError(
                "Excel 파일이 읽기 전용입니다. 열려 있는 파일을 닫고 쓰기 권한을 확인하세요."
            )
        yield workbook
    finally:
        if workbook is not None:
            try:
                _com_call("Workbook.Close", lambda: workbook.Close(SaveChanges=False))
            except Exception:
                logger.exception("Could not close the CellRelay workbook")
        workbook = None


def sheet_names(path: Path) -> list[str]:
    with excel_application() as application:
        try:
            with open_workbook(application, path, read_only=True) as workbook:
                try:
                    return [
                        str(
                            _com_call(
                                "Worksheet.Name",
                                lambda index=index: (
                                    workbook.Worksheets.Item(index).Name
                                ),
                            )
                        )
                        for index in range(
                            1,
                            int(
                                _com_call(
                                    "Worksheets.Count",
                                    lambda: workbook.Worksheets.Count,
                                )
                            )
                            + 1,
                        )
                    ]
                finally:
                    workbook = None
        finally:
            application = None


def read_column(path: Path, sheet: str, column: int, start_row: int) -> tuple:
    """Snapshot a contiguous input column, including zero and False values."""
    values = []
    with excel_application() as application:
        try:
            with open_workbook(application, path, read_only=True) as workbook:
                try:
                    worksheet = _com_call(
                        "Worksheets.Item", lambda: workbook.Worksheets.Item(sheet)
                    )
                    try:
                        values = _column_values(worksheet, column, start_row)
                    finally:
                        worksheet = None
                finally:
                    workbook = None
        finally:
            application = None
    return values


def _column_values(worksheet, column: int, start_row: int) -> tuple:
    values = []
    try:
        if column > int(
            _com_call("Columns.Count", lambda: worksheet.Columns.Count)
        ) or start_row > int(_com_call("Rows.Count", lambda: worksheet.Rows.Count)):
            raise ValueError("선택한 파일 형식의 Excel 범위를 벗어난 셀입니다.")
        used = _com_call("Worksheet.UsedRange", lambda: worksheet.UsedRange)
        last_row = (
            int(_com_call("UsedRange.Row", lambda: used.Row))
            + int(_com_call("UsedRange.Rows.Count", lambda: used.Rows.Count))
            - 1
        )
        used = None
        letter = get_column_letter(column)
        for first in range(start_row, last_row + 1, _READ_BATCH_SIZE):
            last = min(first + _READ_BATCH_SIZE - 1, last_row)
            data = _com_call(
                "Range.Value2.read",
                lambda first=first, last=last: (
                    worksheet.Range(f"{letter}{first}:{letter}{last}").Value2
                ),
            )
            rows = ((data,),) if first == last else data
            for (value,) in rows:
                if value is None or value == "":
                    return tuple(values)
                # Excel exposes numbers as doubles; preserve integer IDs.
                if isinstance(value, float) and value.is_integer():
                    value = int(value)
                values.append(value)
        return tuple(values)
    finally:
        worksheet = None


def save_error_copy(
    path: Path, destination: Path, sheet: str, row: int, header: str, entry: str
) -> None:
    """Save with Excel's original file format and reopen to verify the edit."""
    with excel_application() as application:
        try:
            # Edit an in-memory read-only workbook and save a separate copy.
            # Opening legacy XLS for writing can alter its lock metadata even
            # when it is later closed with SaveChanges=False.
            with open_workbook(application, path, read_only=True) as workbook:
                try:
                    worksheet = _com_call(
                        "Worksheets.Item", lambda: workbook.Worksheets.Item(sheet)
                    )
                    try:
                        expected = _write_error_cells(worksheet, row, header, entry)
                    finally:
                        worksheet = None
                    # SaveCopyAs keeps the original format, including legacy files.
                    _com_call(
                        "Workbook.SaveCopyAs",
                        lambda: workbook.SaveCopyAs(str(destination)),
                    )
                finally:
                    workbook = None
            with open_workbook(application, destination, read_only=True) as check:
                try:
                    worksheet = _com_call(
                        "Worksheets.Item.verify", lambda: check.Worksheets.Item(sheet)
                    )
                    try:
                        for address, text in expected.items():
                            if (
                                _com_call(
                                    "Range.Value2.verify",
                                    lambda address=address: (
                                        worksheet.Range(address).Value2
                                    ),
                                )
                                != text
                            ):
                                raise RuntimeError(
                                    "Excel COM 오류 기록 저장 검증에 실패했습니다."
                                )
                    finally:
                        worksheet = None
                finally:
                    check = None
        finally:
            application = None


def _write_error_cells(worksheet, row: int, header: str, entry: str) -> dict[str, str]:
    try:
        used = _com_call("Worksheet.UsedRange", lambda: worksheet.UsedRange)
        max_col = (
            int(_com_call("UsedRange.Column", lambda: used.Column))
            + int(_com_call("UsedRange.Columns.Count", lambda: used.Columns.Count))
            - 1
        )
        used = None
        data = _com_call(
            "Range.Value2.headers",
            lambda: worksheet.Range(f"A1:{get_column_letter(max_col)}1").Value2,
        )
        headers = (data,) if max_col == 1 else data[0]
        matches = [
            index
            for index, value in enumerate(headers, 1)
            if isinstance(value, str) and value.split("\n", 1)[0] == header
        ]
        if len(matches) > 1:
            raise ValueError("CellRelay 오류 열이 여러 개입니다. 하나만 남겨주세요.")
        column = matches[0] if matches else max_col + 1
        if column > int(_com_call("Columns.Count", lambda: worksheet.Columns.Count)):
            raise ValueError("오류 기록 열을 추가할 공간이 없습니다.")
        if row > int(_com_call("Rows.Count", lambda: worksheet.Rows.Count)):
            raise ValueError("선택한 파일 형식의 마지막 행을 벗어났습니다.")
        letter = get_column_letter(column)
        expected = {}
        if row == 1:
            expected[f"{letter}1"] = (header + "\n" + entry)[:32767]
        else:
            if not matches:
                expected[f"{letter}1"] = header
            expected[f"{letter}{row}"] = entry[:32767]
        for address, text in expected.items():
            cell = _com_call(
                "Worksheet.Range", lambda address=address: worksheet.Range(address)
            )
            try:
                # A failed lambda with cell as a default retains that proxy in
                # its function object even after clearing traceback frame locals.
                _com_call(
                    "Range.NumberFormat.set",
                    partial(setattr, cell, "NumberFormat", "@"),
                )
                _com_call(
                    "Range.Value2.set",
                    partial(setattr, cell, "Value2", text),
                )
            finally:
                cell = None
        return expected
    finally:
        worksheet = None
