"""Single-column, row-by-row Excel reader for CellRelay."""

from __future__ import annotations

import logging
import re
import sys
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.excel import com_excel

_CELL_RE = re.compile(r"^([A-Za-z]{1,3})([1-9][0-9]*)$")
_MAX_EXCEL_COLUMN = 16_384
_MAX_EXCEL_ROW = 1_048_576
logger = logging.getLogger(__name__)


class ExcelReaderError(RuntimeError):
    """Raised for invalid workbook operations or coordinates."""


class ExcelReader:
    """Prefer desktop Excel on Windows, with an OOXML fallback if unavailable."""

    def __init__(self, backend: str = "auto") -> None:
        if backend not in {"auto", "com", "ooxml"}:
            raise ValueError(f"Unknown Excel backend: {backend}")
        self._requested_backend = backend
        self._backend: str | None = None
        self._workbook: Workbook | None = None
        self._worksheet: Worksheet | None = None
        self._path: Path | None = None
        self._sheet_names: list[str] = []
        self._selected_sheet: str | None = None
        self._column_values: tuple | None = None
        self._source_signature: tuple[int, int] | None = None
        self._column = 1
        self._row = 1
        self._start_row = 1
        self._total_items_cache: int | None = None

    @property
    def path(self) -> Path | None:
        return self._path

    @property
    def selected_sheet(self) -> str | None:
        return self._selected_sheet

    @property
    def backend(self) -> str | None:
        return self._backend

    @property
    def sheet_names(self) -> list[str]:
        self._require_loaded()
        return list(self._sheet_names)

    @property
    def current_cell_address(self) -> str:
        return f"{get_column_letter(self._column)}{self._row}"

    @property
    def current_value(self) -> Any:
        if self._backend == "com":
            values = self._require_column_values()
            offset = self._row - self._start_row
            return values[offset] if 0 <= offset < len(values) else None
        worksheet = self._require_worksheet()
        return worksheet.cell(row=self._row, column=self._column).value

    @property
    def current_text(self) -> str | None:
        """Return the current value as text, or None for an empty cell."""
        value = self.current_value
        if value is None or value == "":
            return None
        return str(value)

    def load(self, file_path: str | Path) -> None:
        """Use COM without assuming that the file is an unencrypted ZIP."""
        path = Path(file_path).expanduser().resolve()
        if path.suffix.lower() not in com_excel.SUPPORTED_EXTENSIONS:
            raise ExcelReaderError(".xlsx, .xlsm, .xls, .xlsb 파일만 지원합니다.")
        if not path.is_file():
            raise ExcelReaderError(f"Excel 파일을 찾을 수 없습니다: {path}")
        self.close()
        unavailable = None
        try:
            stat = path.stat()
            self._source_signature = (stat.st_mtime_ns, stat.st_size)
            if self._requested_backend == "com" or (
                self._requested_backend == "auto" and sys.platform == "win32"
            ):
                try:
                    self._sheet_names = com_excel.sheet_names(path)
                    self._backend = "com"
                except com_excel.ExcelComUnavailable as exc:
                    if self._requested_backend == "com":
                        raise
                    unavailable = str(exc)
                    logger.info("Excel COM unavailable; trying OOXML: %s", exc)
            if self._backend is None:
                if path.suffix.lower() not in {".xlsx", ".xlsm"}:
                    raise ExcelReaderError(
                        f"이 파일 형식은 Windows용 Microsoft Excel COM이 필요합니다. {unavailable or ''}"
                    )
                if not com_excel.is_ooxml_workbook(path):
                    raise BadZipFile("Not an OOXML workbook package")
                self._workbook = load_workbook(
                    filename=path, read_only=False, data_only=True
                )
                self._sheet_names = list(self._workbook.sheetnames)
                self._backend = "ooxml"
            stat = path.stat()
            if (stat.st_mtime_ns, stat.st_size) != self._source_signature:
                raise ExcelReaderError(
                    "불러오는 동안 Excel 파일이 변경되었습니다. 다시 시도하세요."
                )
        except BadZipFile as exc:
            self.close()
            raise ExcelReaderError(
                "이 파일은 일반 XLSX ZIP 형식이 아닙니다. "
                "암호화·보안 적용 또는 다른 형식의 파일일 수 있으므로 "
                "Windows용 Microsoft Excel COM으로 열어야 합니다. "
                + (unavailable or "같은 PC의 Excel에서 파일이 열리는지 확인하세요.")
            ) from exc
        except Exception as exc:
            self.close()
            raise ExcelReaderError(f"Excel 파일을 열지 못했습니다: {exc}") from exc
        self._path = path

    def select_sheet(self, sheet_name: str) -> None:
        self._require_loaded()
        if sheet_name not in self._sheet_names:
            raise ExcelReaderError(f"Sheet를 찾을 수 없습니다: {sheet_name}")
        self._selected_sheet = sheet_name
        self._worksheet = self._workbook[sheet_name] if self._workbook else None
        self._column_values = None
        self._total_items_cache = None

    def set_start_cell(self, address: str) -> None:
        """Set the initial coordinate while preserving its column thereafter."""
        self._require_selected_sheet()
        normalized = address.strip().upper()
        match = _CELL_RE.fullmatch(normalized)
        if match is None:
            raise ExcelReaderError(f"올바른 셀 주소가 아닙니다: {address}")
        column_letters, row_text = match.groups()
        column = 0
        for character in column_letters:
            column = column * 26 + (ord(character) - ord("A") + 1)
        row = int(row_text)
        if column > _MAX_EXCEL_COLUMN or row > _MAX_EXCEL_ROW:
            raise ExcelReaderError(f"Excel 범위를 벗어난 셀입니다: {address}")
        self._column = column
        self._row = row
        self._start_row = row
        self._total_items_cache = None
        self._column_values = None
        if self._backend == "com":
            self._require_column_values()

    def advance(self) -> str:
        """Move exactly one row only after the controller records completion."""
        if self._row >= _MAX_EXCEL_ROW:
            raise ExcelReaderError("Excel의 마지막 행에 도달했습니다.")
        self._row += 1
        return self.current_cell_address

    def estimate_total_items(self) -> int:
        """Count contiguous non-empty cells from the configured start cell."""
        if self._total_items_cache is not None:
            return self._total_items_cache
        if self._backend == "com":
            self._total_items_cache = len(self._require_column_values())
            return self._total_items_cache
        worksheet = self._require_worksheet()
        count = 0
        for row in range(self._start_row, min(worksheet.max_row, _MAX_EXCEL_ROW) + 1):
            value = worksheet.cell(row=row, column=self._column).value
            if value is None or value == "":
                break
            count += 1
        self._total_items_cache = count
        return count

    def close(self) -> None:
        if self._workbook is not None:
            self._workbook.close()
        self._workbook = None
        self._worksheet = None
        self._path = None
        self._backend = None
        self._sheet_names = []
        self._selected_sheet = None
        self._column_values = None
        self._source_signature = None
        self._total_items_cache = None

    def _require_loaded(self) -> None:
        if self._path is None:
            raise ExcelReaderError("먼저 Excel 파일을 선택하세요.")

    def _require_selected_sheet(self) -> None:
        if self._selected_sheet is None:
            raise ExcelReaderError("먼저 Sheet를 선택하세요.")

    def _require_column_values(self) -> tuple:
        self._require_selected_sheet()
        if self._column_values is None:
            assert self._path is not None and self._selected_sheet is not None
            stat = self._path.stat()
            if (stat.st_mtime_ns, stat.st_size) != self._source_signature:
                raise ExcelReaderError(
                    "Excel 파일이 외부에서 변경되었습니다. 다시 불러오세요."
                )
            values = com_excel.read_column(
                self._path, self._selected_sheet, self._column, self._start_row
            )
            stat = self._path.stat()
            if (stat.st_mtime_ns, stat.st_size) != self._source_signature:
                raise ExcelReaderError(
                    "읽는 동안 Excel 파일이 변경되었습니다. 다시 불러오세요."
                )
            self._column_values = values
        return self._column_values

    def _require_worksheet(self) -> Worksheet:
        if self._worksheet is None:
            raise ExcelReaderError("먼저 Sheet를 선택하세요.")
        return self._worksheet
