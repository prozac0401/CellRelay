"""Append error cells using the same Excel backend as the reader.

COM handles desktop Excel formats without interpreting them as ZIP files. The
OOXML fallback edits only worksheet XML to preserve other package components.
"""

from __future__ import annotations

import logging
import os
import posixpath
import shutil
import tempfile
import time
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4
from xml.dom import Node, minidom
from zipfile import ZipFile

from openpyxl.utils.cell import (
    column_index_from_string,
    coordinate_from_string,
    get_column_letter,
)
from PySide6.QtCore import QObject, Signal, Slot

from app.excel import com_excel

HEADER = "CellRelay 검색 오류"
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
logger = logging.getLogger(__name__)
# Only sharing/lock violations are transient. Access denied can be permanent.
_TRANSIENT_REPLACE_WINERRORS = {32, 33}
_REPLACE_RETRY_DELAYS = (0.1, 0.2, 0.4, 0.8)


def _exception_code(exc: BaseException, attribute: str):
    seen = set()
    while exc is not None and id(exc) not in seen:
        seen.add(id(exc))
        code = getattr(exc, attribute, None)
        if code is not None:
            return code
        exc = exc.__cause__ or exc.__context__
    return None


@contextmanager
def _write_stage(stage: str, path: Path):
    started = time.monotonic()
    logger.info("Excel write stage=%s event=start path=%s", stage, path)
    try:
        yield
    except Exception as exc:
        winerror = _exception_code(exc, "winerror")
        hresult = _exception_code(exc, "hresult")
        logger.exception(
            "Excel write stage=%s event=failed path=%s elapsed_ms=%.0f "
            "winerror=%s hresult=%s",
            stage,
            path,
            (time.monotonic() - started) * 1000,
            winerror,
            hresult,
        )
        raise
    else:
        logger.info(
            "Excel write stage=%s event=finished path=%s elapsed_ms=%.0f",
            stage,
            path,
            (time.monotonic() - started) * 1000,
        )


def _cleanup_temp(path: Path | None) -> None:
    if path is None:
        return
    started = time.monotonic()
    logger.info("Excel write stage=cleanup event=start path=%s", path)
    try:
        path.unlink(missing_ok=True)
    except OSError as exc:
        # Cleanup must not replace the original failure or a successful commit.
        logger.warning(
            "Excel write stage=cleanup event=failed path=%s elapsed_ms=%.0f winerror=%s",
            path,
            (time.monotonic() - started) * 1000,
            getattr(exc, "winerror", None),
            exc_info=True,
        )
    else:
        logger.info(
            "Excel write stage=cleanup event=finished path=%s elapsed_ms=%.0f",
            path,
            (time.monotonic() - started) * 1000,
        )


def signature(path: Path) -> tuple[int, int]:
    stat = path.stat()
    return stat.st_mtime_ns, stat.st_size


def _elements(node, name):
    return list(node.getElementsByTagNameNS(NS, name))


def _text(node) -> str:
    return "".join(
        child.data if child.nodeType == Node.TEXT_NODE else _text(child)
        for child in node.childNodes
    )


def _cell_text(cell, shared: list[str]) -> str:
    if cell.getAttribute("t") == "s":
        return shared[int(_text(_elements(cell, "v")[0]))]
    if cell.getAttribute("t") == "inlineStr":
        return "".join(_text(t) for t in _elements(cell, "t"))
    values = _elements(cell, "v")
    return _text(values[0]) if values else ""


class ExcelResultWriter:
    """Serial writer owned by ExcelWorker; a backup is created before first edit."""

    def __init__(self) -> None:
        self._backups: dict[Path, Path] = {}

    def write_error(
        self,
        path: Path,
        sheet: str,
        address: str,
        message: str,
        expected_signature: tuple[int, int],
        backend: str = "auto",
    ) -> tuple[tuple[int, int], str]:
        path = path.resolve()
        logger.info(
            "Excel write stage=validate path=%s sheet=%s cell=%s backend=%s",
            path,
            sheet,
            address,
            backend,
        )
        if signature(path) != expected_signature:
            raise RuntimeError(
                "Excel 파일이 외부에서 변경되었습니다. 다시 불러온 뒤 실행하세요."
            )
        if path.suffix.lower() not in com_excel.SUPPORTED_EXTENSIONS:
            raise ValueError(".xlsx, .xlsm, .xls, .xlsb 파일만 지원합니다.")
        if backend not in {"auto", "com", "ooxml"}:
            raise ValueError(f"Unknown Excel backend: {backend}")
        if backend == "com" or (
            backend == "auto"
            and (
                path.suffix.lower() in {".xls", ".xlsb"}
                or not com_excel.is_ooxml_workbook(path)
            )
        ):
            return self._write_com_error(
                path, sheet, address, message, expected_signature
            )
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ValueError("이 파일 형식은 Excel COM으로 저장해야 합니다.")
        _, row_number = coordinate_from_string(address)
        temp_path: Path | None = None
        try:
            with _write_stage("prepare_ooxml", path), ZipFile(path) as source:
                workbook = minidom.parseString(source.read("xl/workbook.xml"))
                target = [
                    s
                    for s in _elements(workbook, "sheet")
                    if s.getAttribute("name") == sheet
                ]
                if len(target) != 1:
                    raise ValueError("결과를 기록할 Sheet를 찾지 못했습니다.")
                rel_id = target[0].getAttributeNS(REL_NS, "id")
                relationships = minidom.parseString(
                    source.read("xl/_rels/workbook.xml.rels")
                )
                rels = [
                    r
                    for r in relationships.documentElement.childNodes
                    if r.nodeType == Node.ELEMENT_NODE
                    and r.getAttribute("Id") == rel_id
                ]
                if len(rels) != 1 or rels[0].getAttribute("TargetMode") == "External":
                    raise ValueError("Sheet 경로를 확인하지 못했습니다.")
                relative = rels[0].getAttribute("Target")
                sheet_path = posixpath.normpath(
                    relative.lstrip("/")
                    if relative.startswith("/")
                    else posixpath.join("xl", relative)
                )
                if (
                    not sheet_path.startswith("xl/")
                    or sheet_path not in source.namelist()
                ):
                    raise ValueError("잘못된 Sheet 경로입니다.")
                doc = minidom.parseString(source.read(sheet_path))
                shared = []
                if "xl/sharedStrings.xml" in source.namelist():
                    strings = minidom.parseString(source.read("xl/sharedStrings.xml"))
                    shared = [
                        "".join(_text(t) for t in _elements(si, "t"))
                        for si in _elements(strings, "si")
                    ]
                cells = _elements(doc, "c")
                result_cols = []
                max_col = 0
                max_row = 1
                for cell in cells:
                    col, row = coordinate_from_string(cell.getAttribute("r"))
                    index = column_index_from_string(col)
                    max_col, max_row = max(max_col, index), max(max_row, row)
                    if (
                        row == 1
                        and _cell_text(cell, shared).split("\n", 1)[0] == HEADER
                    ):
                        result_cols.append(index)
                if len(result_cols) > 1:
                    raise ValueError(
                        "CellRelay 오류 열이 여러 개입니다. 하나만 남겨주세요."
                    )
                # Include merged/declared ranges so an appended column cannot
                # land inside existing formatted or merged data.
                for dimension in _elements(doc, "dimension") + _elements(
                    doc, "mergeCell"
                ):
                    ref = dimension.getAttribute("ref")
                    if ref:
                        col, row = coordinate_from_string(ref.split(":")[-1])
                        max_col, max_row = (
                            max(max_col, column_index_from_string(col)),
                            max(max_row, row),
                        )
                result_col = result_cols[0] if result_cols else max_col + 1
                if result_col > 16384:
                    raise ValueError("오류 기록 열을 추가할 공간이 없습니다.")
                data = _elements(doc, "sheetData")[0]
                prefix = (
                    (doc.documentElement.prefix + ":")
                    if doc.documentElement.prefix
                    else ""
                )

                def element(name):
                    return doc.createElementNS(NS, prefix + name)

                def write_cell(row_number, text):
                    rows = [
                        r
                        for r in _elements(data, "row")
                        if int(r.getAttribute("r")) == row_number
                    ]
                    if rows:
                        row = rows[0]
                    else:
                        row = element("row")
                        row.setAttribute("r", str(row_number))
                        following = next(
                            (
                                r
                                for r in _elements(data, "row")
                                if int(r.getAttribute("r")) > row_number
                            ),
                            None,
                        )
                        data.insertBefore(row, following)
                    coord = f"{get_column_letter(result_col)}{row_number}"
                    current = next(
                        (
                            c
                            for c in _elements(row, "c")
                            if c.getAttribute("r") == coord
                        ),
                        None,
                    )
                    if current is None:
                        current = element("c")
                        current.setAttribute("r", coord)
                        following = next(
                            (
                                c
                                for c in _elements(row, "c")
                                if column_index_from_string(
                                    coordinate_from_string(c.getAttribute("r"))[0]
                                )
                                > result_col
                            ),
                            None,
                        )
                        row.insertBefore(current, following)
                    for child in list(current.childNodes):
                        current.removeChild(child)
                    current.setAttribute("t", "inlineStr")
                    inline, content = element("is"), element("t")
                    content.setAttribute("xml:space", "preserve")
                    # Playwright errors can contain terminal control codes,
                    # which are not legal XML 1.0 characters.
                    clean = "".join(
                        ch
                        for ch in text
                        if ch in "\t\n\r"
                        or 0x20 <= ord(ch) <= 0xD7FF
                        or 0xE000 <= ord(ch) <= 0xFFFD
                        or 0x10000 <= ord(ch) <= 0x10FFFF
                    )
                    content.appendChild(doc.createTextNode(clean[:32767]))
                    inline.appendChild(content)
                    current.appendChild(inline)

                stamp = (
                    datetime.now(timezone.utc)
                    .astimezone()
                    .isoformat(timespec="seconds")
                )
                entry = f"{stamp} | {address} | {message}"
                # When input starts on row 1, keep both the column identifier
                # and the row-1 error in that cell (do not insert a data row).
                if row_number == 1:
                    write_cell(1, HEADER + "\n" + entry)
                elif not result_cols:
                    write_cell(1, HEADER)
                if row_number != 1:
                    write_cell(row_number, entry)
                dimensions = _elements(doc, "dimension")
                if dimensions:
                    dimensions[0].setAttribute(
                        "ref",
                        f"A1:{get_column_letter(max(max_col, result_col))}{max(max_row, row_number)}",
                    )
                with tempfile.NamedTemporaryFile(
                    prefix=".cellrelay-",
                    suffix=path.suffix,
                    dir=path.parent,
                    delete=False,
                ) as tmp:
                    temp_path = Path(tmp.name)
                with ZipFile(temp_path, "w") as output:
                    for info in source.infolist():
                        output.writestr(
                            info,
                            doc.toxml(encoding="utf-8")
                            if info.filename == sheet_path
                            else source.read(info.filename),
                        )
                with (
                    _write_stage("verify_ooxml", temp_path),
                    ZipFile(temp_path) as check,
                ):
                    if check.testzip() is not None:
                        raise RuntimeError("결과 Excel 파일 검증에 실패했습니다.")
                    minidom.parseString(check.read(sheet_path))
            return self._replace_with_backup(path, temp_path, expected_signature)
        finally:
            _cleanup_temp(temp_path)

    def _write_com_error(self, path, sheet, address, message, expected_signature):
        _, row = coordinate_from_string(address)
        stamp = datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")
        entry = f"{stamp} | {address} | {message}"
        entry = "".join(
            ch
            for ch in entry
            if ch in "\t\n\r"
            or 0x20 <= ord(ch) <= 0xD7FF
            or 0xE000 <= ord(ch) <= 0xFFFD
            or 0x10000 <= ord(ch) <= 0x10FFFF
        )
        temp_path = None
        try:
            with tempfile.NamedTemporaryFile(
                prefix=".cellrelay-", suffix=path.suffix, dir=path.parent, delete=False
            ) as tmp:
                temp_path = Path(tmp.name)
            with _write_stage("save_com_copy", path):
                com_excel.save_error_copy(path, temp_path, sheet, row, HEADER, entry)
            return self._replace_with_backup(path, temp_path, expected_signature)
        finally:
            _cleanup_temp(temp_path)

    def _replace_with_backup(self, path, temp_path, expected_signature):
        if signature(path) != expected_signature:
            raise RuntimeError(
                "저장 중 Excel 파일이 변경되어 원본을 덮어쓰지 않았습니다."
            )
        if path not in self._backups:
            backup = path.with_name(
                f"{path.stem}.cellrelay-backup-{uuid4().hex[:12]}{path.suffix}"
            )
            with _write_stage("backup", path):
                shutil.copy2(path, backup)
            self._backups[path] = backup
        with _write_stage("replace", path):
            for attempt in range(len(_REPLACE_RETRY_DELAYS) + 1):
                if signature(path) != expected_signature:
                    raise RuntimeError(
                        "백업 또는 저장 재시도 중 Excel 파일이 변경되어 원본을 덮어쓰지 않았습니다."
                    )
                try:
                    os.replace(temp_path, path)
                    break
                except OSError as exc:
                    winerror = getattr(exc, "winerror", None)
                    if winerror in _TRANSIENT_REPLACE_WINERRORS and attempt < len(
                        _REPLACE_RETRY_DELAYS
                    ):
                        delay = _REPLACE_RETRY_DELAYS[attempt]
                        logger.warning(
                            "Excel write stage=replace event=retry path=%s "
                            "attempt=%d max_attempts=%d wait_ms=%.0f winerror=%s",
                            path,
                            attempt + 1,
                            len(_REPLACE_RETRY_DELAYS) + 1,
                            delay * 1000,
                            winerror,
                        )
                        time.sleep(delay)
                        continue
                    if isinstance(exc, PermissionError) or (
                        winerror in _TRANSIENT_REPLACE_WINERRORS
                    ):
                        raise PermissionError(
                            "Excel 파일이 열려 있거나 읽기 전용입니다. "
                            "파일을 닫고 쓰기 권한을 확인하세요."
                        ) from exc
                    raise
        return signature(path), str(self._backups[path])


class ExcelWorker(QObject):
    """Execute disk writes off the UI thread and acknowledge every result."""

    finished = Signal(int, bool, object, str)

    def __init__(self) -> None:
        super().__init__()
        self.writer = ExcelResultWriter()

    @Slot(int, object)
    def write_error(self, run_id: int, request: dict) -> None:
        started = time.monotonic()
        metadata = request if isinstance(request, dict) else {}
        context = (
            run_id,
            metadata.get("path"),
            metadata.get("sheet"),
            metadata.get("address"),
            metadata.get("backend", "auto"),
        )
        logger.info(
            "Excel write event=start run_id=%s path=%s sheet=%s cell=%s backend=%s",
            *context,
        )
        try:
            sig, backup = self.writer.write_error(**request)
            logger.info(
                "Excel write event=finished run_id=%s path=%s sheet=%s cell=%s "
                "backend=%s elapsed_ms=%.0f",
                *context,
                (time.monotonic() - started) * 1000,
            )
            self.finished.emit(run_id, True, sig, backup)
        except Exception as exc:
            winerror = _exception_code(exc, "winerror")
            hresult = _exception_code(exc, "hresult")
            logger.exception(
                "Excel write event=failed run_id=%s path=%s sheet=%s cell=%s "
                "backend=%s elapsed_ms=%.0f winerror=%s hresult=%s",
                *context,
                (time.monotonic() - started) * 1000,
                winerror,
                hresult,
            )
            self.finished.emit(run_id, False, None, str(exc))
