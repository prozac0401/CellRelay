import time
from functools import partial

import pytest
from openpyxl import Workbook, load_workbook
from PySide6.QtCore import QCoreApplication

from app.config.settings import AWS_LOGIN_URL
from app.core.controller import CellRelayController
from app.core.state import AppState
from app.excel.excel_reader import ExcelReader


def drain_until(predicate, timeout=5):
    deadline = time.monotonic() + timeout
    while not predicate() and time.monotonic() < deadline:
        QCoreApplication.processEvents()
        time.sleep(0.01)
    assert predicate()


@pytest.fixture
def controller(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "app.core.controller.ExcelReader", partial(ExcelReader, backend="ooxml")
    )
    app = QCoreApplication.instance() or QCoreApplication([])
    book = Workbook()
    book.active.title = "Data"
    book.active["B5"] = "first"
    book.active["B6"] = "second"
    path = tmp_path / "users.xlsx"
    book.save(path)
    c = CellRelayController(tmp_path)
    # Keep production QThreads and queued stop/save acknowledgements, but do
    # not send any real browser commands in these state-machine tests.
    c._aws_assign_command.disconnect()
    c._browser_open = True
    c._browser_url = AWS_LOGIN_URL
    c._aws_page_ready = True
    c._confirmed_training_url = "https://skillbuilder.aws/admin/organization/modality/curriculum/training/one?orgId=org"
    c.start_job(str(path), "Data", "B5", AWS_LOGIN_URL, "textarea", "aws_skill_builder")
    assert c.state == AppState.INPUTTING
    yield c
    c.shutdown()
    drain_until(
        lambda: not c._worker_thread.isRunning() and not c._excel_thread.isRunning()
    )
    app.processEvents()


def test_stop_blocks_restart_until_queued_ack(controller):
    c = controller
    old_run = c._run_id
    c.stop_job()
    assert c.state == AppState.STOPPING
    c.start_job(
        c.settings.excel_file,
        "Data",
        "B5",
        AWS_LOGIN_URL,
        "textarea",
        "aws_skill_builder",
    )
    assert c._run_id == old_run
    assert c._worker._stop_event.is_set()
    drain_until(lambda: c.state == AppState.READY)
    assert c._run_id != old_run
    assert c.progress.current_cell == "B5"


def test_receipt_is_saved_while_paused_and_advanced_once_on_resume(controller):
    c = controller
    c.pause_job()
    c._on_assignment_completed(c._run_id)
    c._on_assignment_completed(c._run_id)
    assert c.state == AppState.PAUSED
    assert c.progress.current_cell == "B5"
    assert c.progress.processed_count == 1
    assert c._store.load_progress().last_completed_cell == "B5"
    assert c._store.load_progress().phase == "AWS_ASSIGNMENT_SUBMITTED"
    assert c._store.load_progress().terminal_outcome == "AWS_ASSIGNMENT_SUBMITTED"
    c.resume_job()
    assert c.progress.current_cell == "B6"
    assert c.progress.processed_count == 1


def test_queued_success_after_stop_is_not_discarded(controller):
    c = controller
    run_id = c._run_id
    c.stop_job()
    c._on_assignment_completed(run_id)
    drain_until(lambda: c.state == AppState.READY)
    assert c.progress.current_cell == "B6"
    assert c._store.load_progress().last_completed_cell == "B5"


def test_skip_waits_for_excel_ack_before_advancing(controller):
    c = controller
    c._on_assignment_skipped(c._run_id, "NOT_FOUND: no matching user")
    assert c.progress.current_cell == "B5"
    assert c._excel_pending
    drain_until(lambda: c.progress.current_cell == "B6")
    assert c.progress.skipped_count == 1
    book = load_workbook(c.settings.excel_file)
    assert "NOT_FOUND" in book["Data"]["C5"].value
    book.close()


def test_save_failure_stops_without_skipping(controller):
    c = controller
    c._file_signature = (0, 0)
    c._on_assignment_skipped(c._run_id, "NOT_FOUND")
    drain_until(lambda: c.state == AppState.ERROR)
    assert c.progress.current_cell == "B5"
    assert c.progress.skipped_count == 0
    assert "저장 실패" in c.progress.last_message
    assert "원래 작업 결과: NOT_FOUND" in c.progress.last_message
    assert "NOT_FOUND" in c._store.load_progress().last_error
    assert not c._excel_wait_timer.isActive()


def test_uncertain_browser_error_recorded_but_not_skipped(controller):
    c = controller
    c._on_worker_failure("aws_assign_user", c._run_id, "SEARCH_TIMEOUT", "")
    drain_until(lambda: c.state == AppState.ERROR)
    assert c.progress.current_cell == "B5"
    assert c.progress.skipped_count == 0
    book = load_workbook(c.settings.excel_file)
    assert "SEARCH_TIMEOUT" in book["Data"]["C5"].value
    book.close()


def test_confirmation_does_not_replace_login_url(controller):
    c = controller
    c.stop_job()
    drain_until(lambda: c.state == AppState.READY)
    target = c._confirmed_training_url
    c._on_aws_page_tested(True, "confirmed", target)
    assert c.settings.url == AWS_LOGIN_URL
    assert c._confirmed_training_url == target


def test_next_row_persistence_failure_does_not_advance_twice(controller, monkeypatch):
    c = controller
    original = c._store.save_progress

    def save(progress):
        if progress.phase == "READY_TO_INPUT":
            raise OSError("disk full")
        original(progress)

    monkeypatch.setattr(c._store, "save_progress", save)
    c._on_assignment_completed(c._run_id)
    drain_until(lambda: c.state == AppState.ERROR)
    assert c.progress.current_cell == "B6"
    assert c.progress.processed_count == 1
    assert c._store.load_progress().last_completed_cell == "B5"


def test_shutdown_waits_for_pending_excel_write(controller):
    c = controller
    c._on_assignment_skipped(c._run_id, "NOT_FOUND")
    c.shutdown()
    assert not c._shutdown_started
    drain_until(lambda: c._shutdown_started)
    assert c.progress.skipped_count == 1
    assert c.progress.current_cell == "B6"


@pytest.mark.parametrize("control", ["running", "paused", "stopping"])
def test_slow_excel_report_keeps_pending_cell_until_matching_ack(controller, control):
    c = controller
    c._write_excel_command.disconnect()
    requests = []
    c._write_excel_command.connect(lambda run_id, request: requests.append(request))
    run_id = c._run_id
    try:
        c._on_assignment_skipped(run_id, "NOT_FOUND: test result")
        c._excel_pending_since = time.monotonic() - 16
        if control == "stopping":
            c.stop_job()
            drain_until(lambda: c._stop_ack)
        elif control == "paused":
            c.pause_job()
        expected_state = {
            "running": AppState.INPUTTING,
            "paused": AppState.PAUSED,
            "stopping": AppState.STOPPING,
        }[control]
        c._report_excel_wait()
        c._queue_current_value()
        assert c.state == expected_state
        assert c._excel_pending and c._excel_wait_timer.isActive()
        assert c.progress.current_cell == "B5"
        assert c.progress.processed_count == 0
        assert len(requests) == 1
        assert "응답을 기다리고" in c.progress.last_message
        assert "16초" in c.progress.last_message

        # A stale worker reply must not dismiss the active write or its notice.
        c._on_excel_written(run_id - 1, True, c._file_signature, "")
        assert c._excel_pending and c._excel_wait_timer.isActive()
        c._on_excel_written(run_id, True, c._file_signature, "")
        assert not c._excel_pending and not c._excel_wait_timer.isActive()
        if control == "paused":
            assert c.progress.current_cell == "B5"
            assert "기록을 완료" in c.progress.last_message
            c.resume_job()
        assert c.progress.current_cell == "B6"
        assert c.progress.skipped_count == 1
        message = c.progress.last_message
        c._report_excel_wait()
        assert c.progress.last_message == message
    finally:
        if c._excel_pending:
            c._on_excel_written(run_id, False, None, "test cleanup")


def test_submitted_rows_advance_without_excel_errors_and_finish_with_review_notice(
    controller,
):
    c = controller
    before = c._file_signature
    c._on_assignment_completed(c._run_id)
    assert c.progress.current_cell == "B6"
    c._on_assignment_completed(c._run_id)
    assert c.state == AppState.COMPLETED
    assert c.progress.current_cell == "B7"
    assert c.progress.processed_count == 2
    assert c.progress.skipped_count == 0
    assert not c._excel_pending
    assert c._file_signature == before
    assert "관리자가 최종 등록 명단" in c.progress.last_message
    book = load_workbook(c.settings.excel_file)
    assert book["Data"].max_column == 2
    book.close()
