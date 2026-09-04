from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from app.excel.excel_reader import ExcelReader, ExcelReaderError
from app.excel.result_writer import HEADER, ExcelResultWriter, signature


def _make_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["B5"] = "first"
    sheet["B6"] = 42
    sheet["B7"] = ""
    sheet["B8"] = "must not be reached"
    workbook.create_sheet("Other")
    workbook.save(path)


def test_reads_down_one_column_and_stops_at_first_empty_cell(tmp_path: Path) -> None:
    workbook_path = tmp_path / "relay.xlsx"
    _make_workbook(workbook_path)
    reader = ExcelReader()

    reader.load(workbook_path)
    assert reader.sheet_names == ["Data", "Other"]
    reader.select_sheet("Data")
    reader.set_start_cell("b5")

    assert reader.current_cell_address == "B5"
    assert reader.current_text == "first"
    assert reader.estimate_total_items() == 2
    assert reader.advance() == "B6"
    assert reader.current_text == "42"
    assert reader.advance() == "B7"
    assert reader.current_text is None


@pytest.mark.parametrize("address", ["", "B0", "A1:B2", "XFE1", "A1048577"])
def test_rejects_invalid_or_out_of_range_cell_addresses(
    tmp_path: Path, address: str
) -> None:
    workbook_path = tmp_path / "relay.xlsx"
    _make_workbook(workbook_path)
    reader = ExcelReader()
    reader.load(workbook_path)
    reader.select_sheet("Data")

    with pytest.raises(ExcelReaderError):
        reader.set_start_cell(address)


def test_rejects_non_xlsx_file(tmp_path: Path) -> None:
    path = tmp_path / "relay.xls"
    path.write_bytes(b"not a workbook")
    with pytest.raises(ExcelReaderError, match="xlsx"):
        ExcelReader().load(path)


def test_error_column_without_changing_values_or_font_style(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "relay.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["B5"] = "not-assigned"
    sheet["B5"].font = Font(name="Arial", size=14, bold=True)
    sheet["B6"] = "next-user"
    sheet["C5"] = "=1+1"
    workbook.save(workbook_path)

    reader = ExcelReader()
    reader.load(workbook_path)
    reader.select_sheet("Data")
    reader.set_start_cell("B5")
    _, backup = ExcelResultWriter().write_error(
        workbook_path, "Data", "B5", "NOT_FOUND", signature(workbook_path)
    )
    assert Path(backup).is_file()

    # The running reader keeps its values and can continue to the next row.
    assert reader.current_text == "not-assigned"
    assert reader.advance() == "B6"
    assert reader.current_text == "next-user"

    saved = load_workbook(workbook_path, data_only=False)
    try:
        marked = saved["Data"]["B5"]
        assert marked.value == "not-assigned"
        assert marked.font.bold
        assert marked.font.name == "Arial"
        assert marked.font.sz == 14
        assert marked.font.color is None
        assert saved["Data"]["D1"].value == HEADER
        assert "NOT_FOUND" in saved["Data"]["D5"].value
        assert saved["Data"]["B6"].value == "next-user"
        assert saved["Data"]["C5"].value == "=1+1"
    finally:
        saved.close()
