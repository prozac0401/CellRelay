import logging
from pathlib import Path
from unittest.mock import Mock
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from app.excel import result_writer
from app.excel.result_writer import HEADER, ExcelResultWriter, signature


def workbook(path):
    book = Workbook()
    book.active.title = "Data"
    book.active["B1"] = "first"
    book.active["B2"] = "second"
    book.active["C2"] = "=1+1"
    book.save(path)
    book.close()
    # A cached formula value and an opaque package part must survive reporting.
    with ZipFile(path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    doc = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
    cell = doc.find(f'.//{ns}c[@r="C2"]')
    value = cell.find(f"{ns}v")
    if value is None:
        value = ET.SubElement(cell, f"{ns}v")
    value.text = "2"
    parts["xl/worksheets/sheet1.xml"] = ET.tostring(doc)
    parts["customXml/opaque.xml"] = b"<preserve>opaque extension</preserve>"
    with ZipFile(path, "w") as dest:
        for name, data in parts.items():
            dest.writestr(name, data)
    return parts


def test_reuses_column_preserves_first_row_formulas_cache_and_package(tmp_path):
    path = tmp_path / "users.xlsx"
    before = workbook(path)
    writer = ExcelResultWriter()
    sig, backup = writer.write_error(path, "Data", "B1", "NOT_FOUND", signature(path))
    _, same_backup = writer.write_error(path, "Data", "B2", "MULTIPLE_MATCHES", sig)
    assert backup == same_backup
    with ZipFile(path) as result, ZipFile(backup) as original:
        for name, data in before.items():
            assert original.read(name) == data
            if name != "xl/worksheets/sheet1.xml":
                assert result.read(name) == data
    saved = load_workbook(path, data_only=True)
    assert saved.active.max_column == 4
    assert saved.active["D1"].value.startswith(HEADER + "\n")
    assert "NOT_FOUND" in saved.active["D1"].value
    assert "MULTIPLE_MATCHES" in saved.active["D2"].value
    assert saved.active["C2"].value == 2
    assert saved.active["B1"].value == "first"
    saved.close()
    formulas = load_workbook(path, data_only=False)
    assert formulas.active["C2"].value == "=1+1"
    formulas.close()


def test_external_change_refused(tmp_path):
    path = tmp_path / "users.xlsx"
    workbook(path)
    with pytest.raises(RuntimeError, match="외부에서 변경"):
        ExcelResultWriter().write_error(path, "Data", "B2", "NOT_FOUND", (0, 0))
    assert not list(tmp_path.glob("*.cellrelay-backup-*.xlsx"))


def test_locked_destination_leaves_original_and_backup_intact(tmp_path, monkeypatch):
    path = tmp_path / "users.xlsx"
    workbook(path)
    original = path.read_bytes()

    def locked(*args):
        raise PermissionError("workbook is open in Excel")

    monkeypatch.setattr("app.excel.result_writer.os.replace", locked)
    with pytest.raises(PermissionError):
        ExcelResultWriter().write_error(
            path, "Data", "B2", "NOT_FOUND", signature(path)
        )
    assert path.read_bytes() == original
    backup = next(tmp_path.glob("*.cellrelay-backup-*.xlsx"))
    assert backup.read_bytes() == original
    assert not list(tmp_path.glob(".cellrelay-*.xlsx"))


def windows_error(code):
    error = PermissionError(13, "simulated Windows file access error")
    if code is not None:
        error.winerror = code
    return error


@pytest.mark.parametrize("winerror", [32, 33])
def test_transient_replace_lock_recovers_without_rebuilding_copy(
    tmp_path, monkeypatch, winerror, caplog
):
    path = tmp_path / "users.xlsb"
    path.write_bytes(b"original workbook")
    original = path.read_bytes()
    real_replace = result_writer.os.replace
    attempts = []

    def replace(source, destination):
        attempts.append((source, destination))
        if len(attempts) == 1:
            raise windows_error(winerror)
        return real_replace(source, destination)

    def save_copy(source, destination, *args):
        destination.write_bytes(b"verified workbook with error report")

    save = Mock(side_effect=save_copy)
    sleep = Mock()
    monkeypatch.setattr(result_writer.os, "replace", replace)
    monkeypatch.setattr(result_writer.time, "sleep", sleep)
    monkeypatch.setattr(result_writer.com_excel, "save_error_copy", save)
    with caplog.at_level(logging.INFO, logger=result_writer.__name__):
        sig, backup = ExcelResultWriter().write_error(
            path, "Data", "B5", "NOT_FOUND", signature(path), backend="com"
        )
    assert sig == signature(path)
    assert path.read_bytes() == b"verified workbook with error report"
    assert Path(backup).read_bytes() == original
    assert len(attempts) == 2 and attempts[0] == attempts[1]
    save.assert_called_once()
    sleep.assert_called_once_with(result_writer._REPLACE_RETRY_DELAYS[0])
    assert len(list(tmp_path.glob("*.cellrelay-backup-*"))) == 1
    assert not list(tmp_path.glob(".cellrelay-*"))
    assert "event=retry" in caplog.text and f"winerror={winerror}" in caplog.text
    assert "stage=replace event=finished" in caplog.text


@pytest.mark.parametrize("winerror", [None, 5])
def test_permanent_access_error_is_not_retried(tmp_path, monkeypatch, winerror):
    path = tmp_path / "users.xlsx"
    workbook(path)
    original = path.read_bytes()
    error = windows_error(winerror)
    replace, sleep = Mock(side_effect=error), Mock()
    monkeypatch.setattr(result_writer.os, "replace", replace)
    monkeypatch.setattr(result_writer.time, "sleep", sleep)
    with pytest.raises(PermissionError) as caught:
        ExcelResultWriter().write_error(
            path, "Data", "B2", "NOT_FOUND", signature(path)
        )
    replace.assert_called_once()
    sleep.assert_not_called()
    assert caught.value.__cause__ is error
    assert path.read_bytes() == original
    assert not list(tmp_path.glob(".cellrelay-*"))


def test_transient_lock_has_bounded_retries_and_preserves_original(
    tmp_path, monkeypatch
):
    path = tmp_path / "users.xlsx"
    workbook(path)
    original = path.read_bytes()
    error = windows_error(32)
    replace, sleep = Mock(side_effect=error), Mock()
    monkeypatch.setattr(result_writer.os, "replace", replace)
    monkeypatch.setattr(result_writer.time, "sleep", sleep)
    with pytest.raises(PermissionError) as caught:
        ExcelResultWriter().write_error(
            path, "Data", "B2", "NOT_FOUND", signature(path)
        )
    assert replace.call_count == len(result_writer._REPLACE_RETRY_DELAYS) + 1
    assert [call.args[0] for call in sleep.call_args_list] == list(
        result_writer._REPLACE_RETRY_DELAYS
    )
    assert sum(result_writer._REPLACE_RETRY_DELAYS) <= 2
    assert caught.value.__cause__ is error
    assert path.read_bytes() == original
    backups = list(tmp_path.glob("*.cellrelay-backup-*"))
    assert len(backups) == 1 and backups[0].read_bytes() == original
    assert not list(tmp_path.glob(".cellrelay-*"))


def test_external_change_during_retry_is_never_overwritten(tmp_path, monkeypatch):
    path = tmp_path / "users.xlsx"
    workbook(path)
    original = path.read_bytes()
    replace = Mock(side_effect=windows_error(32))
    sleep = Mock(side_effect=lambda delay: path.write_bytes(b"external edit"))
    monkeypatch.setattr(result_writer.os, "replace", replace)
    monkeypatch.setattr(result_writer.time, "sleep", sleep)
    with pytest.raises(RuntimeError, match="변경되어 원본을 덮어쓰지"):
        ExcelResultWriter().write_error(
            path, "Data", "B2", "NOT_FOUND", signature(path)
        )
    replace.assert_called_once()
    sleep.assert_called_once()
    assert path.read_bytes() == b"external edit"
    backups = list(tmp_path.glob("*.cellrelay-backup-*"))
    assert len(backups) == 1 and backups[0].read_bytes() == original
    assert not list(tmp_path.glob(".cellrelay-*"))


@pytest.mark.parametrize("backend", ["com", "ooxml"])
@pytest.mark.parametrize("save_succeeds", [False, True])
def test_cleanup_error_preserves_original_save_outcome(
    tmp_path, monkeypatch, caplog, backend, save_succeeds
):
    path = tmp_path / "users.xlsx"
    workbook(path)
    original = path.read_bytes()
    writer = ExcelResultWriter()
    primary = RuntimeError("original save failure")
    real_unlink = Path.unlink

    def cleanup_locked(candidate, *args, **kwargs):
        if candidate.name.startswith(".cellrelay-"):
            raise windows_error(32)
        return real_unlink(candidate, *args, **kwargs)

    def save_copy(source, destination, *args):
        if not save_succeeds:
            raise primary
        destination.write_bytes(b"verified result")

    with monkeypatch.context() as scoped:
        scoped.setattr(Path, "unlink", cleanup_locked)
        scoped.setattr(result_writer.com_excel, "save_error_copy", save_copy)
        if not save_succeeds and backend == "ooxml":
            scoped.setattr(writer, "_replace_with_backup", Mock(side_effect=primary))
        if save_succeeds:
            sig, backup = writer.write_error(
                path, "Data", "B2", "NOT_FOUND", signature(path), backend=backend
            )
            assert sig == signature(path)
            assert Path(backup).read_bytes() == original
            assert path.read_bytes() != original
        else:
            with pytest.raises(RuntimeError) as caught:
                writer.write_error(
                    path, "Data", "B2", "NOT_FOUND", signature(path), backend=backend
                )
            assert caught.value is primary
            assert path.read_bytes() == original
    assert "stage=cleanup event=failed" in caplog.text
    for leftover in tmp_path.glob(".cellrelay-*"):
        leftover.unlink()
